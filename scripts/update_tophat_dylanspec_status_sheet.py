#!/usr/bin/env python3
"""Build a shared status sheet for Dylan-spectrum top-hat reruns."""

from __future__ import annotations

import argparse
import copy
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


DEFAULT_EVENTS = [
    "050525A",
    "050922C",
    "090424",
    "090618",
    "111228A",
    "130612A",
    "131030A",
    "140506A",
    "171010A",
    "210905A",
    "220101A",
]

EVENT_ALIASES = {
    "161031A": "160131A",
}


def now_utc_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--vegas-dir", default="/Users/jkeohane/GRBs/VegasJetFit")
    parser.add_argument(
        "--drive-root",
        default="/Users/jkeohane/My Drive (jwkeohane@gmail.com)/VegasGRBruns",
    )
    parser.add_argument("--owner-subdir", default="jkeohane")
    parser.add_argument(
        "--source-run-tag",
        default="theta1p0_thesis_short_kmin10_seeded_v1",
    )
    parser.add_argument(
        "--run-tag",
        default="theta1p0_thesis_short_kmin10_seeded_vegasv201_dylanspec_v1",
    )
    parser.add_argument("--events", nargs="+")
    parser.add_argument(
        "--output-csv",
        default="/Users/jkeohane/My Drive (jwkeohane@gmail.com)/VegasGRBruns/Tophat_DylanSpectrum_Run_Status.csv",
    )
    parser.add_argument(
        "--output-xlsx",
        default="/Users/jkeohane/My Drive (jwkeohane@gmail.com)/VegasGRBruns/Tophat_DylanSpectrum_Run_Status.xlsx",
    )
    parser.add_argument(
        "--assignment",
        action="append",
        default=[],
        help="Optional host:event1,event2,... mapping.",
    )
    parser.add_argument(
        "--tracking-sheet",
        help="Optional CSV/XLSX GRB tracking sheet to update in place.",
    )
    parser.add_argument(
        "--tracking-sheet-name",
        help="Optional XLSX sheet name. Defaults to the first sheet.",
    )
    parser.add_argument(
        "--tracking-event-column",
        default="GRB (Priority Order)",
        help="Column in the tracking sheet containing the event IDs.",
    )
    parser.add_argument(
        "--tracking-status-column",
        default="Restart2000 Status",
        help="Column name to create/update with the current run status.",
    )
    return parser.parse_args()


def load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text())
    except Exception:
        return {}


def to_float_or_none(value: Any) -> float | None:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(out):
        return out
    return out


def valid_minimized_payload(payload: dict[str, Any]) -> bool:
    success = payload.get("success")
    nmap = payload.get("nmap")
    return success is True and isinstance(nmap, (int, float)) and math.isfinite(float(nmap))


def parse_assignments(values: list[str]) -> tuple[dict[str, str], dict[str, int]]:
    host_map: dict[str, str] = {}
    queue_index: dict[str, int] = {}
    for value in values:
        if ":" not in value:
            continue
        host, raw_events = value.split(":", 1)
        events = [item for item in raw_events.split(",") if item]
        for idx, event in enumerate(events, start=1):
            host_map[event] = host
            queue_index[event] = idx
    return host_map, queue_index


def normalize_event_name(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = EVENT_ALIASES.get(text, text)
    if re.fullmatch(r"\d{6}[A-Z]?", text):
        return text
    return None


def load_tracking_sheet(path: Path, sheet_name: str | None) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path)
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        target_sheet = 0 if not sheet_name else sheet_name
        return pd.read_excel(path, sheet_name=target_sheet, engine="openpyxl")
    raise ValueError(f"Unsupported tracking sheet format: {path}")


def ordered_events_from_tracking_sheet(
    path: Path,
    *,
    sheet_name: str | None,
    event_column: str,
) -> list[str]:
    df = load_tracking_sheet(path, sheet_name)
    source_column = event_column if event_column in df.columns else df.columns[0]
    events: list[str] = []
    seen: set[str] = set()
    for value in df[source_column].tolist():
        event = normalize_event_name(value)
        if not event or event in seen:
            continue
        seen.add(event)
        events.append(event)
    return events


def status_label(
    *,
    assigned_host: str,
    queue_index: int | None,
    has_best: bool,
    has_chain: bool,
    has_summary: bool,
    has_minimized: bool,
    has_share_manifest: bool,
    has_share_minimized: bool,
) -> str:
    if has_share_manifest and has_share_minimized:
        return "synced"
    if has_minimized:
        return "minimized_pending_sync"
    if has_best and has_chain and has_summary:
        return "complete_pending_minimize"
    if has_best or has_chain:
        return "running_or_partial"
    if assigned_host:
        if queue_index is not None:
            return f"queued on {assigned_host} (#{queue_index})"
        return f"queued on {assigned_host}"
    return "not_started"


def write_xlsx(df: pd.DataFrame, out_xlsx: Path) -> None:
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="status")
        ws = writer.sheets["status"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                text = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(text))
            ws.column_dimensions[col_letter].width = min(max(10, max_len + 2), 42)


def update_tracking_csv(
    path: Path,
    *,
    event_column: str,
    status_column: str,
    status_by_event: dict[str, str],
) -> None:
    df = pd.read_csv(path)
    source_column = event_column if event_column in df.columns else df.columns[0]
    values: list[str] = []
    for raw_value in df[source_column].tolist():
        event = normalize_event_name(raw_value)
        values.append(status_by_event.get(event, "") if event else "")
    df[status_column] = values
    df.to_csv(path, index=False)


def update_tracking_xlsx(
    path: Path,
    *,
    sheet_name: str | None,
    event_column: str,
    status_column: str,
    status_by_event: dict[str, str],
) -> None:
    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]

    header_row = 1
    header_cells = list(ws[header_row])
    header_map = {str(cell.value).strip(): cell.column for cell in header_cells if cell.value is not None}
    event_col = header_map.get(event_column, header_cells[0].column)
    status_col = header_map.get(status_column)

    if status_col is None:
        status_col = ws.max_column + 1
        header_source_col = max(1, status_col - 1)
        source_header = ws.cell(row=header_row, column=header_source_col)
        target_header = ws.cell(row=header_row, column=status_col)
        if source_header.has_style:
            target_header._style = copy.copy(source_header._style)
        if source_header.font:
            target_header.font = copy.copy(source_header.font)
        if source_header.fill:
            target_header.fill = copy.copy(source_header.fill)
        if source_header.border:
            target_header.border = copy.copy(source_header.border)
        if source_header.alignment:
            target_header.alignment = copy.copy(source_header.alignment)
        if source_header.number_format:
            target_header.number_format = source_header.number_format
        if source_header.protection:
            target_header.protection = copy.copy(source_header.protection)
        target_header.value = status_column

    for row_idx in range(header_row + 1, ws.max_row + 1):
        raw_value = ws.cell(row=row_idx, column=event_col).value
        event = normalize_event_name(raw_value)
        target_cell = ws.cell(row=row_idx, column=status_col)
        if status_col > 1:
            source_cell = ws.cell(row=row_idx, column=status_col - 1)
            if source_cell.has_style:
                target_cell._style = copy.copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.alignment:
                target_cell.alignment = copy.copy(source_cell.alignment)
        target_cell.value = status_by_event.get(event, "") if event else ""

    wb.save(path)


def update_tracking_sheet(
    path: Path,
    *,
    sheet_name: str | None,
    event_column: str,
    status_column: str,
    status_by_event: dict[str, str],
) -> None:
    if path.suffix.lower() == ".csv":
        update_tracking_csv(
            path,
            event_column=event_column,
            status_column=status_column,
            status_by_event=status_by_event,
        )
        return
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        update_tracking_xlsx(
            path,
            sheet_name=sheet_name,
            event_column=event_column,
            status_column=status_column,
            status_by_event=status_by_event,
        )
        return
    raise ValueError(f"Unsupported tracking sheet format: {path}")


def main() -> int:
    args = parse_args()

    vegas_dir = Path(args.vegas_dir)
    results_root = vegas_dir / "jetfit" / "results"
    drive_root = Path(args.drive_root)
    timestamp = now_utc_iso()
    host_map, queue_index = parse_assignments(args.assignment)
    tracking_sheet = Path(args.tracking_sheet) if args.tracking_sheet else None

    if args.events:
        events = args.events
    elif tracking_sheet:
        events = ordered_events_from_tracking_sheet(
            tracking_sheet,
            sheet_name=args.tracking_sheet_name,
            event_column=args.tracking_event_column,
        )
    else:
        events = DEFAULT_EVENTS

    rows: list[dict[str, Any]] = []
    for event in events:
        source_name = f"{event}_powerlaw_tophat_{args.source_run_tag}"
        run_name = f"{event}_powerlaw_tophat_{args.run_tag}"
        run_dir = results_root / run_name
        source_dir = results_root / source_name

        best_fit_path = run_dir / "best_fit.json"
        chain_path = run_dir / "chain.npz"
        summary_path = run_dir / "summary.csv"
        checkpoint_path = run_dir / "pt_resume_state.npz"
        minimized_path = run_dir / "minimized" / "minimized.json"

        share_dir = drive_root / event / args.owner_subdir / run_name
        share_manifest = share_dir / "sync_manifest.txt"
        share_minimized = share_dir / "minimized" / "minimized.json"

        best_fit = load_json(best_fit_path)
        minimized = load_json(minimized_path)
        share_minimized_payload = load_json(share_minimized)
        has_valid_minimized = minimized_path.exists() and valid_minimized_payload(minimized)
        has_valid_share_minimized = share_minimized.exists() and valid_minimized_payload(share_minimized_payload)

        row = {
            "timestamp_utc": timestamp,
            "event": event,
            "assigned_host": host_map.get(event, ""),
            "queue_index": queue_index.get(event),
            "source_run_dir": str(source_dir),
            "run_dir": str(run_dir),
            "share_dir": str(share_dir),
            "status": status_label(
                assigned_host=host_map.get(event, ""),
                queue_index=queue_index.get(event),
                has_best=best_fit_path.exists(),
                has_chain=chain_path.exists(),
                has_summary=summary_path.exists(),
                has_minimized=has_valid_minimized,
                has_share_manifest=share_manifest.exists(),
                has_share_minimized=has_valid_share_minimized,
            ),
            "has_best_fit": best_fit_path.exists(),
            "has_chain": chain_path.exists(),
            "has_summary": summary_path.exists(),
            "has_checkpoint": checkpoint_path.exists(),
            "has_minimized": has_valid_minimized,
            "has_share_manifest": share_manifest.exists(),
            "has_share_minimized": has_valid_share_minimized,
            "nmap_best_fit": to_float_or_none(best_fit.get("nmap")),
            "nmap_minimized": to_float_or_none(minimized.get("nmap")),
            "minimize_success": minimized.get("success"),
            "minimize_message": minimized.get("message"),
        }
        rows.append(row)

    df = pd.DataFrame(rows)
    csv_path = Path(args.output_csv)
    xlsx_path = Path(args.output_xlsx)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(csv_path, index=False)
    write_xlsx(df, xlsx_path)
    if tracking_sheet:
        status_by_event = dict(zip(df["event"], df["status"]))
        update_tracking_sheet(
            tracking_sheet,
            sheet_name=args.tracking_sheet_name,
            event_column=args.tracking_event_column,
            status_column=args.tracking_status_column,
            status_by_event=status_by_event,
        )
    print(csv_path)
    print(xlsx_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
